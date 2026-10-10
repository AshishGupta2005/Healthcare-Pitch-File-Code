"""Shared helpers for the healthcare pitch deck tools."""
import re
import pandas as pd

MISSING = {"", "nan", "none", "not verified", "undisclosed", "not publicly verified", "not located / verify"}
URL_RE = re.compile(r"^https?://\S+$")
AMOUNT_RE = re.compile(r"^\$\d[\d,]*(\.\d+)?[KMB]$")
ALLOWED_STAGES = {"Pre-Seed", "Seed", "Crowdfunding", "Series A", "Series B", "Series C", "Series D",
                  "Series E", "Late Stage", "Not verified", "Undisclosed"}
PROMO_WORDS = ["leading", "innovative", "revolutionary", "cutting-edge", "cutting edge", "best-in-class",
               "world-class", "disruptive", "disrupting", "game-changing", "groundbreaking", "state-of-the-art",
               "seamless", "next-generation", "unicorn", "visionary", "pioneering", "transformative", "unparalleled"]

def is_missing(v):
    return str(v).strip().lower() in MISSING

def load_decks(path, sheet="Healthcare Decks"):
    """Read the deck sheet. Where a cell shows text but carries a hyperlink (e.g. the Website column shows the
    company name and links to its site), the link target is used as the value."""
    from openpyxl import load_workbook
    df = pd.read_excel(path, sheet_name=sheet, dtype=object)
    df.columns = [str(c).strip() for c in df.columns]
    df = df[[c for c in df.columns if not c.startswith("Unnamed")]]
    if "Name" not in df.columns:
        raise SystemExit(f"Sheet '{sheet}' in {path} has no 'Name' column. Found: {list(df.columns)}")
    ws = load_workbook(path)[sheet]
    header = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    for i in range(len(df)):
        xl_row = i + 2
        if str(ws.cell(xl_row, header["Name"]).value).strip() != str(df.iloc[i]["Name"]).strip():
            continue
        for col, idx in header.items():
            if col not in df.columns:
                continue
            link = ws.cell(xl_row, idx).hyperlink
            cur = str(df.iloc[i][col]).strip()
            if link is not None and link.target and not URL_RE.match(cur):
                df.iat[i, df.columns.get_loc(col)] = link.target
    return df.fillna("")

def stage_mismatch(summary, stage):
    m = re.search(r"Series ([A-E]) deck", str(summary))
    if m and str(stage).strip() and str(stage).strip() != f"Series {m.group(1)}":
        return f"Series {m.group(1)}"
    return None

def amount_conflict(amount_text, usd_m_cell):
    """Return the numeric column value when it disagrees with the Amount Raised text, else None."""
    a = usd_millions(amount_text)
    try:
        b = float(usd_m_cell)
    except (TypeError, ValueError):
        return None
    return b if a is not None and abs(a - b) > 0.01 else None

def usd_millions(text):
    m = re.match(r"^\$([\d,]+(?:\.\d+)?)([KMB])$", str(text).strip())
    if not m:
        return None
    return float(m.group(1).replace(",", "")) * {"K": 0.001, "M": 1, "B": 1000}[m.group(2)]

def promo_hits(text):
    t = str(text).lower()
    return [w for w in PROMO_WORDS if re.search(r"\b" + re.escape(w) + r"\b", t)]
