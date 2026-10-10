#!/usr/bin/env python3
"""
Generate company summaries from the database workbook, as plain-text files and one Word document.

Input : DATABASE.xlsx
          sheet 'Healthcare Decks' (required), sheet 'Asset Tracker' (optional)
        company_descriptions.csv (optional: Name, Sector, Description) for the business sentence
Optional deck_notes.csv (Name, Topic, Text, Source) adds a 'Deck content' section for the companies listed.
Output: <out>/TXT/<Company>_summary.txt (one per company), <out>/Company_Summaries_All.txt,
        <out>/Company_Summaries.docx

Every statement comes from the workbook. Figures are shown as recorded. Where the workbook disagrees with
itself (for example, the Amount Raised text differs from the Amount (USD M) column), the summary says so
and does not choose between the two. Website links are read from cell hyperlinks.

Usage:
  python build_summaries.py DATABASE.xlsx --out ./Company_Summaries
  python build_summaries.py DATABASE.xlsx --descriptions company_descriptions.csv --out ./Company_Summaries
"""
import argparse, datetime, os, re
import pandas as pd
from common import load_decks, is_missing, URL_RE, promo_hits, stage_mismatch, amount_conflict

NV = "Not verified"
STD_NOTE = ("Funding details come from secondary directories. Check them against the original deck or a "
            "funding announcement before external use.")

def val(r, k):
    v = str(r.get(k, "")).strip()
    if v.endswith(".0"): v = v[:-2]
    if v.lower() == "undisclosed": return "Undisclosed"
    return NV if is_missing(v) else v

def cap(t):
    t = t.strip()
    return (t[0].upper() + t[1:]) if t else t

def split_summary(text):
    parts = [p.strip() for p in str(text).split(";") if p.strip()]
    return (parts[0].rstrip(".") if parts else ""), [cap(p.rstrip(".")) + "." for p in parts[1:]]

def funding_sentence(r):
    amt, st, yr = val(r, "Amount Raised"), val(r, "Stage"), val(r, "Year")
    a = {"Undisclosed": "an undisclosed amount", NV: "no verified amount"}.get(amt, amt)
    s = {"Undisclosed": "a round with an undisclosed stage", NV: "a round whose stage is not verified"}.get(st, f"a {st} round")
    y = {"Undisclosed": "in an undisclosed year", NV: "in a year that is not verified"}.get(yr, f"in {yr}")
    return f"The workbook records {a} for {s} {y}."

def business_sentence(r, desc):
    name = r["Name"]
    d = desc.get(name)
    base, _ = split_summary(r.get("Company / Deck Summary", ""))
    if d and d.get("Description"):
        return f"{name} {d['Description']}."
    return f"Business description recorded in the workbook: {base}." if base else "No business description is recorded."

def build_record(r, desc, assets, notes_by_name=None):
    name = r["Name"]
    sector = val(r, "Industry / Sub-Sector") if "Industry / Sub-Sector" in r else (desc.get(name, {}).get("Sector") or NV)
    if sector == NV and desc.get(name, {}).get("Sector"): sector = desc[name]["Sector"]
    _, notes = split_summary(r.get("Company / Deck Summary", ""))
    a = assets.get(name, {})
    biz = business_sentence(r, desc) + (f" It is listed under {sector}." if sector != NV else "")
    paras = [biz, funding_sentence(r),
             "Amount, year and stage describe the deck listing in the workbook; check whether the amount is the round "
             "size or a lifetime total before comparing companies."]
    data_notes = []
    c = amount_conflict(r.get("Amount Raised", ""), r.get("Amount (USD M)", ""))
    if c is not None:
        data_notes.append(f"The workbook's Amount (USD M) column shows {c:g}, which differs from the Amount Raised text "
                          f"({val(r, 'Amount Raised')}). Confirm which figure belongs to the round shown in the Stage and Year fields.")
    mm = stage_mismatch(r.get("Company / Deck Summary", ""), r.get("Stage", ""))
    if mm:
        data_notes.append(f"The workbook summary refers to a {mm} deck, while the Stage column shows {val(r, 'Stage')}.")
    flag = str(r.get("Review Flag", "")).strip()
    if flag: data_notes.append(f"Workbook review note: {flag}")
    data_notes += [f"Workbook note: {n}" for n in notes]
    return dict(
        name=name, sector=sector, website=val(r, "Website"), amount=val(r, "Amount Raised"), year=val(r, "Year"),
        stage=val(r, "Stage"), usd=str(r.get("Amount (USD M)", "")).strip().removesuffix(".0"),
        completeness=val(r, "Data Completeness"), source=val(r, "Source Page"),
        hosted=a.get("Deck Hosted Via", NV), deck_file=a.get("Deck File (PDF/PPT)", "Not locally downloaded"),
        next_step=a.get("Next Step", NV), paras=paras, data_notes=data_notes,
        deck=(notes_by_name or {}).get(name, []))

def to_txt(x):
    L = [f"COMPANY SUMMARY: {x['name']}", "=" * (17 + len(x["name"])), "",
         f"Company:   {x['name']}", "Industry:  Healthcare", f"Sector:    {x['sector']}", f"Website:   {x['website']}", "",
         "FUNDING RECORD (as recorded in the workbook)",
         f"Amount Raised:           {x['amount']}", f"Year:                    {x['year']}", f"Stage:                   {x['stage']}"]
    if x["usd"]: L.append(f"Amount (USD M) column:   {x['usd']}")
    L += ["", "BUSINESS SUMMARY"] + [p for p in x["paras"][:1]] + ["", "FUNDING CONTEXT"] + x["paras"][1:]
    if x["deck"]:
        L += ["", "DECK CONTENT (from the deck's source page text; the slides themselves were not readable)"]
        for topic, text, src in x["deck"]: L.append(f"- {topic}: {text}")
        L.append(f"  Source: {x['deck'][0][2]}")
    L += [
          "", "DECK AND SOURCES",
          f"Source page:      {x['source']}", f"Deck hosting:     {x['hosted']}", f"Deck file:        {x['deck_file']}",
          "Slide text:       Not extracted; this file is a summary only.", f"Next step:        {x['next_step']}", "",
          "DATA STATUS", f"Completeness:     {x['completeness']}"]
    for n in x["data_notes"]: L.append(f"- {n}")
    L += [f"- {STD_NOTE}", ""]
    return "\n".join(L)

def write_docx(recs, path, src_name):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    NAVY = RGBColor(0x1F, 0x38, 0x64)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial"); st.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 12)):
        h = doc.styles[name]; h.font.name = "Arial"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = NAVY
        rf = h.element.rPr.rFonts
        for k in [k for k in rf.attrib if "Theme" in k]: del rf.attrib[k]
        for k in ("ascii", "hAnsi", "eastAsia", "cs"): rf.set(qn("w:" + k), "Arial")
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(16); doc.styles["Heading 1"].paragraph_format.space_after = Pt(6)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(10); doc.styles["Heading 2"].paragraph_format.space_after = Pt(3)

    def shade(cell, hex_):
        tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_); tcPr.append(sh)
    def set_w(cell, cm): cell.width = Cm(cm)
    def para_border_bottom(p):
        pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "1F3864")): bt.set(qn("w:" + k), v)
        b.append(bt); pPr.append(b)
    def cell_text(cell, text, bold=False, color=None, size=9, align=None):
        cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        if align: p.alignment = align

    # footer page numbers
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Healthcare Pitch Deck Research  |  Page "); r.font.size = Pt(8)
    for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        r2 = fp.add_run(); r2.font.size = Pt(8)
        if t:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
        r2._r.append(e)

    # title block
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(40)
    t = doc.add_paragraph("Healthcare Pitch Deck Research", style="Title")
    sub = doc.add_paragraph(); rr = sub.add_run("Company Summaries"); rr.font.size = Pt(16); rr.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    para_border_bottom(sub)
    meta = doc.add_paragraph(f"{len(recs)} companies  |  Source workbook: {src_name}  |  {datetime.date.today():%B %d, %Y}")
    meta.runs[0].font.size = Pt(9); meta.runs[0].font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    doc.add_heading("About this document", level=1)
    for line in [
        "Each section summarises one company listed in the research workbook: what the company does, the funding record "
        "attached to its deck listing, where the deck can be found, and the status of the data.",
        "All figures are shown as recorded in the workbook. Amount, year and stage describe the deck listing and are not "
        "verified against original decks. Where the workbook disagrees with itself, the section for that company says so.",
        "No deck files have been downloaded and no slide-by-slide text has been extracted. These summaries are not "
        "transcripts of the decks. Funding details come from secondary directories; check them against the original "
        "deck or a funding announcement before external use."]:
        doc.add_paragraph(line)
    n_conf = sum(any("Amount (USD M) column shows" in n for n in x["data_notes"]) for x in recs)
    if n_conf:
        doc.add_paragraph(f"Data note: in {n_conf} of {len(recs)} rows the Amount Raised text differs from the Amount (USD M) "
                          "column of the same row. These rows are marked in their sections and should be reconciled before the "
                          "figures are used externally.")

    doc.add_heading("Overview of companies", level=1)
    cols = ["Company", "Stage", "Year", "Amount Raised", "Completeness"]; widths = [5.6, 2.8, 1.8, 3.4, 3.4]
    tb = doc.add_table(rows=1, cols=len(cols)); tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cols):
        cell_text(tb.rows[0].cells[i], c, bold=True, color=RGBColor(255, 255, 255)); shade(tb.rows[0].cells[i], "1F3864"); set_w(tb.rows[0].cells[i], widths[i])
    trPr = tb.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for k, x in enumerate(recs):
        row = tb.add_row().cells
        for i, v in enumerate([x["name"], x["stage"], x["year"], x["amount"], x["completeness"]]):
            cell_text(row[i], str(v)); set_w(row[i], widths[i])
            if k % 2: shade(row[i], "F2F5FA")

    doc.add_page_break()
    doc.add_heading("Company summaries", level=1)
    for x in recs:
        h = doc.add_heading(x["name"], level=2); h.paragraph_format.keep_with_next = True
        tb = doc.add_table(rows=0, cols=4); tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        pairs = [("Sector", x["sector"], "Stage", x["stage"]), ("Amount Raised", x["amount"], "Year", x["year"]),
                 ("Website", x["website"], "Completeness", x["completeness"])]
        for a, b, c, d in pairs:
            cells = tb.add_row().cells
            for i, (v, lab) in enumerate(((a, True), (b, False), (c, True), (d, False))):
                cell_text(cells[i], v, bold=lab, size=9); set_w(cells[i], (3.0, 6.0, 3.0, 5.0)[i])
                if lab: shade(cells[i], "E8EDF5")
            row_tr = cells[0]._tc.getparent(); cs = OxmlElement("w:cantSplit"); row_tr.get_or_add_trPr().append(cs)
        lab = doc.add_paragraph(); lab.paragraph_format.space_before = Pt(6); lab.paragraph_format.keep_with_next = True
        rr = lab.add_run("Business summary"); rr.bold = True; rr.font.color.rgb = NAVY
        doc.add_paragraph(x["paras"][0])
        lab = doc.add_paragraph(); lab.paragraph_format.keep_with_next = True
        rr = lab.add_run("Funding context"); rr.bold = True; rr.font.color.rgb = NAVY
        doc.add_paragraph(x["paras"][1])
        if x["deck"]:
            lab = doc.add_paragraph(); lab.paragraph_format.keep_with_next = True
            rr = lab.add_run("Deck content"); rr.bold = True; rr.font.color.rgb = NAVY
            ip = doc.add_paragraph("From the text of the deck's source page. The slides themselves were not readable."); ip.runs[0].italic = True
            for topic, text, src in x["deck"]:
                bp = doc.add_paragraph(style="List Bullet"); b = bp.add_run(topic + ": "); b.bold = True; bp.add_run(text)
                bp.paragraph_format.space_after = Pt(2); bp.paragraph_format.keep_together = True
            sp = doc.add_paragraph(f"Source: {x['deck'][0][2]}"); sp.runs[0].font.size = Pt(8)
        lab = doc.add_paragraph(); lab.paragraph_format.keep_with_next = True
        rr = lab.add_run("Deck and sources"); rr.bold = True; rr.font.color.rgb = NAVY
        doc.add_paragraph(f"Source page: {x['source']}. Deck hosting: {x['hosted']}. Deck file: {x['deck_file']}. "
                          f"Slide text has not been extracted. Next step: {x['next_step']}")
        lab = doc.add_paragraph(); lab.paragraph_format.keep_with_next = True
        rr = lab.add_run("Data status"); rr.bold = True; rr.font.color.rgb = NAVY
        for n in x["data_notes"]:
            bp = doc.add_paragraph(style="List Bullet"); bp.add_run(n); bp.paragraph_format.space_after = Pt(2)
            bp.paragraph_format.keep_together = True
        if not x["data_notes"]:
            doc.add_paragraph("No open data notes for this entry.")
        end = doc.add_paragraph(); para_border_bottom(end); end.paragraph_format.space_after = Pt(4)
    z = doc.settings.element.find(qn("w:zoom"))
    if z is not None and z.get(qn("w:percent")) is None: z.set(qn("w:percent"), "100")
    doc.save(path)

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook"); ap.add_argument("--descriptions", default=None)
    ap.add_argument("--out", default="Company_Summaries"); ap.add_argument("--sheet", default="Healthcare Decks")
    ap.add_argument("--deck-notes", default=None, help="CSV with Name, Topic, Text, Source (default: deck_notes.csv next to this script)")
    ap.add_argument("--no-docx", action="store_true")
    a = ap.parse_args()
    df = load_decks(a.workbook, a.sheet)
    desc = {}
    csvp = a.descriptions or os.path.join(os.path.dirname(os.path.abspath(__file__)), "company_descriptions.csv")
    if os.path.exists(csvp):
        desc = pd.read_csv(csvp, dtype=str).fillna("").set_index("Name").to_dict("index")
    assets = {}
    try:
        at = pd.read_excel(a.workbook, sheet_name="Asset Tracker", dtype=object).fillna("")
        assets = at.set_index("Name").to_dict("index")
    except Exception:
        pass
    notes_by_name = {}
    dn = a.deck_notes or os.path.join(os.path.dirname(os.path.abspath(__file__)), "deck_notes.csv")
    if os.path.exists(dn):
        for _, n in pd.read_csv(dn, dtype=str).fillna("").iterrows():
            notes_by_name.setdefault(n["Name"], []).append((n["Topic"], n["Text"], n["Source"]))
    os.makedirs(os.path.join(a.out, "TXT"), exist_ok=True)
    recs = [build_record(r, desc, assets, notes_by_name) for _, r in df.sort_values("Name", key=lambda s: s.str.lower()).iterrows()]
    allt, hits = ["COMPANY SUMMARIES", "", f"{len(recs)} companies. Source workbook: {os.path.basename(a.workbook)}", ""], 0
    for x in recs:
        t = to_txt(x)
        for w in promo_hits(t): print(f"Promotional wording in {x['name']}: {w}"); hits += 1
        fn = re.sub(r"[^A-Za-z0-9]+", "_", x["name"]).strip("_") + "_summary.txt"
        open(os.path.join(a.out, "TXT", fn), "w", encoding="utf-8").write(t)
        allt += [t, "-" * 72, ""]
    open(os.path.join(a.out, "Company_Summaries_All.txt"), "w", encoding="utf-8").write("\n".join(allt))
    if not a.no_docx:
        write_docx(recs, os.path.join(a.out, "Company_Summaries.docx"), os.path.basename(a.workbook))
    print(f"{len(recs)} summaries written to {a.out} ({hits} wording warnings)")

if __name__ == "__main__":
    main()
